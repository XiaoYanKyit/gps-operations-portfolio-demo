from django.contrib.auth.decorators import login_required
from django.db import connections
from django.http import HttpResponse
from django.shortcuts import render
from openpyxl import load_workbook, Workbook
from .models import GPSResetLog, OutletRemovalLog


@login_required
def home(request):
    return render(
        request,
        "partners/dashboard.html"
    )



@login_required
def gps_cancel(request):
    success_message = None
    error_message = None
    results = None

    # Portfolio demo: simulate GPS resets without SAP/PHM access.

    if request.method == "POST" and request.POST.get("action") == "single_reset":
        card_code = request.POST.get("customer_code", "").strip().upper()

        if not card_code:
            error_message = "Please enter a CardCode."
        else:
            ship_to_code = f"{card_code}_S1"

            GPSResetLog.objects.create(
                user=request.user,
                card_code=card_code,
                ship_to_code=ship_to_code,
                reset_type="SINGLE",
                status="SUCCESS",
                message="DEMO ONLY: GPS reset simulated; no real system was changed.",
            )

            success_message = (
                f"DEMO ONLY: Simulated GPS reset for {card_code}. "
                "No real system was changed."
            )

    elif request.method == "POST" and request.POST.get("action") == "bulk_reset":
        excel_file = request.FILES.get("excel_file")

        if not excel_file:
            error_message = "Please select an Excel file."
        elif not excel_file.name.lower().endswith(".xlsx"):
            error_message = "Please upload an .xlsx Excel file."
        else:
            workbook = None

            try:
                workbook = load_workbook(
                    excel_file,
                    read_only=True,
                    data_only=True,
                )
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    raise ValueError("The Excel file is empty.")

                headers = [
                    str(header).strip() if header is not None else ""
                    for header in headers
                ]

                required_columns = ["CardCode", "Ship to Code"]
                for column in required_columns:
                    if column not in headers:
                        raise ValueError(f"Missing required column: {column}")

                card_index = headers.index("CardCode")
                ship_index = headers.index("Ship to Code")
                outlets = []

                for row_number, row in enumerate(rows, start=2):
                    card_code = (
                        str(row[card_index]).strip().upper()
                        if row[card_index] is not None else ""
                    )
                    ship_to_code = (
                        str(row[ship_index]).strip().upper()
                        if row[ship_index] is not None else ""
                    )

                    if not card_code and not ship_to_code:
                        continue
                    if not card_code:
                        raise ValueError(f"Row {row_number}: CardCode is empty.")
                    if not ship_to_code:
                        raise ValueError(f"Row {row_number}: Ship to Code is empty.")

                    outlets.append({
                        "card_code": card_code,
                        "ship_to_code": ship_to_code,
                    })

                if not outlets:
                    raise ValueError("No outlet records were found.")
                if len(outlets) > 500:
                    raise ValueError("Maximum 500 outlets per upload.")

                results = []

                for outlet in outlets:
                    card_code = outlet["card_code"]
                    ship_to_code = outlet["ship_to_code"]
                    message = (
                        "DEMO ONLY: GPS reset simulated; "
                        "no real system was changed."
                    )

                    GPSResetLog.objects.create(
                        user=request.user,
                        card_code=card_code,
                        ship_to_code=ship_to_code,
                        reset_type="BULK",
                        status="SUCCESS",
                        message=message,
                    )

                    results.append({
                        "card_code": card_code,
                        "ship_to_code": ship_to_code,
                        "status": "SUCCESS",
                        "message": message,
                    })

                success_message = (
                    f"DEMO ONLY: {len(results)} GPS reset(s) simulated. "
                    "No real system was changed."
                )

            except Exception as e:
                error_message = str(e)

            finally:
                if workbook is not None:
                    workbook.close()

    return render(
        request,
        "partners/gps_cancel.html",
        {
            "success_message": success_message,
            "error_message": error_message,
            "results": results,
        },
    )


@login_required
def activity(request):

    if not request.user.is_staff:
        return render(
            request,
            "partners/access_denied.html",
            status=403
        )

    # =========================================================
    # SEARCH
    # =========================================================

    search = request.GET.get(
        "search",
        ""
    ).strip()

    # =========================================================
    # GET GPS RESET LOGS
    # =========================================================

    gps_logs = GPSResetLog.objects.select_related(
        "user"
    )

    # =========================================================
    # GET OUTLET REMOVAL LOGS
    # =========================================================

    removal_logs = OutletRemovalLog.objects.select_related(
        "user"
    )

    # =========================================================
    # SEARCH
    # =========================================================

    if search:

        gps_logs = gps_logs.filter(
            card_code__icontains=search
        )

        removal_logs = removal_logs.filter(
            code__icontains=search
        ) | removal_logs.filter(
            customer_code__icontains=search
        )

    # =========================================================
    # COMBINE ACTIVITY LOGS
    # =========================================================

    activities = []

    # ---------------------------------------------------------
    # GPS RESET ACTIVITIES
    # ---------------------------------------------------------

    for log in gps_logs:

        activities.append({
            "created_at": log.created_at,
            "user": log.user,

            "action": "GPS RESET",

            "code": log.card_code,

            "reference": log.ship_to_code,

            "type": log.reset_type,

            "status": log.status,

            "message": log.message,
        })

    # ---------------------------------------------------------
    # OUTLET REMOVAL ACTIVITIES
    # ---------------------------------------------------------

    for log in removal_logs:

        activities.append({
            "created_at": log.created_at,
            "user": log.user,

            "action": "REMOVE OUTLET",

            "code": log.code,

            "reference": log.customer_code,

            "type": log.removal_type,

            "status": log.status,

            "message": log.message,
        })

    # =========================================================
    # SORT NEWEST FIRST
    # =========================================================

    activities.sort(
        key=lambda x: x["created_at"],
        reverse=True
    )

    # =========================================================
    # SUMMARY
    # =========================================================

    total_activities = (
        GPSResetLog.objects.count()
        +
        OutletRemovalLog.objects.count()
    )

    successful_activities = (
        GPSResetLog.objects.filter(
            status="SUCCESS"
        ).count()
        +
        OutletRemovalLog.objects.filter(
            status="SUCCESS"
        ).count()
    )

    failed_activities = (
        GPSResetLog.objects.filter(
            status="FAILED"
        ).count()
        +
        OutletRemovalLog.objects.filter(
            status="FAILED"
        ).count()
    )

    # =========================================================
    # ACTIVE USERS
    # =========================================================

    gps_users = set(
        GPSResetLog.objects
        .exclude(user=None)
        .values_list(
            "user_id",
            flat=True
        )
    )

    removal_users = set(
        OutletRemovalLog.objects
        .exclude(user=None)
        .values_list(
            "user_id",
            flat=True
        )
    )

    active_users = len(
        gps_users | removal_users
    )

    # =========================================================
    # RENDER
    # =========================================================

    return render(
        request,
        "partners/activity.html",
        {
            "activities": activities,

            "total_activities": total_activities,

            "successful_activities": (
                successful_activities
            ),

            "failed_activities": (
                failed_activities
            ),

            "active_users": active_users,

            "search": search,
        }
    )


@login_required
def remove_outlet(request):
    preview_results = None
    removal_results = None
    error_message = None
    success_message = None

    # Portfolio demo only: never connect to PHM or delete real records.

    if request.method == "POST" and request.POST.get("action") == "preview":
        excel_file = request.FILES.get("excel_file")

        if not excel_file:
            error_message = "Please select an Excel file."
        elif not excel_file.name.lower().endswith(".xlsx"):
            error_message = "Only .xlsx Excel files are allowed."
        else:
            workbook = None

            try:
                workbook = load_workbook(
                    excel_file,
                    read_only=True,
                    data_only=True,
                )
                worksheet = workbook.active
                rows = worksheet.iter_rows(values_only=True)
                headers = next(rows, None)

                if not headers:
                    raise ValueError("The Excel file is empty.")

                headers = [
                    str(header).strip() if header is not None else ""
                    for header in headers
                ]

                required_columns = ["Code", "U_CustomerCode"]
                for column in required_columns:
                    if column not in headers:
                        raise ValueError(f"Missing required column: {column}")

                code_index = headers.index("Code")
                customer_index = headers.index("U_CustomerCode")
                records = []

                for row_number, row in enumerate(rows, start=2):
                    code = (
                        str(row[code_index]).strip().upper()
                        if row[code_index] is not None else ""
                    )
                    customer_code = (
                        str(row[customer_index]).strip().upper()
                        if row[customer_index] is not None else ""
                    )

                    if not code and not customer_code:
                        continue
                    if not code:
                        raise ValueError(f"Row {row_number}: Code is empty.")
                    if not customer_code:
                        raise ValueError(
                            f"Row {row_number}: U_CustomerCode is empty."
                        )

                    records.append({
                        "row": row_number,
                        "code": code,
                        "customer_code": customer_code,
                    })

                if not records:
                    raise ValueError("No records were found in the Excel file.")
                if len(records) > 500:
                    raise ValueError("Maximum 500 outlets per upload.")

                # No database lookup: show an explicit demo-only preview.
                preview_results = [
                    {
                        "row": record["row"],
                        "code": record["code"],
                        "customer_code": record["customer_code"],
                        "status": "DEMO",
                        "message": "Demo preview only; PHM record was not checked.",
                    }
                    for record in records
                ]

            except Exception as e:
                error_message = str(e)

            finally:
                if workbook is not None:
                    workbook.close()

    elif request.method == "POST" and request.POST.get("action") == "remove":
        codes = request.POST.getlist("code")
        customer_codes = request.POST.getlist("customer_code")

        if not codes or len(codes) != len(customer_codes):
            error_message = "No valid records were selected."
        elif len(codes) > 500:
            error_message = "Maximum 500 outlets can be processed per upload."
        else:
            removal_results = []

            for code, customer_code in zip(codes, customer_codes):
                code = code.strip().upper()
                customer_code = customer_code.strip().upper()

                if not code or not customer_code:
                    continue

                message = (
                    "DEMO ONLY: Outlet removal simulated. "
                    "No PHM or company data was changed."
                )

                OutletRemovalLog.objects.create(
                    user=request.user,
                    code=code,
                    customer_code=customer_code,
                    removal_type="BULK",
                    status="SUCCESS",
                    message=message,
                )

                removal_results.append({
                    "code": code,
                    "customer_code": customer_code,
                    "status": "SUCCESS",
                    "message": message,
                })

            success_message = (
                f"DEMO ONLY: {len(removal_results)} outlet removal(s) "
                "simulated. No real records were deleted."
            )

    return render(
        request,
        "partners/remove_outlet.html",
        {
            "preview_results": preview_results,
            "removal_results": removal_results,
            "error_message": error_message,
            "success_message": success_message,
        },
    )

@login_required
def export_activity_logs(request):
    activities = []

    gps_logs = GPSResetLog.objects.all()
    outlet_logs = OutletRemovalLog.objects.all()

    for log in gps_logs:
        activities.append({
            "date_time": log.created_at,
            "user": log.user.username if log.user else "",
            "action": "GPS RESET",
            "code": log.card_code,
            "reference": log.ship_to_code,
            "type": log.reset_type,
            "status": log.status,
            "message": log.message,
        })

    for log in outlet_logs:
        activities.append({
            "date_time": log.created_at,
            "user": log.user.username if log.user else "",
            "action": "REMOVE OUTLET",
            "code": log.code,
            "reference": log.customer_code,
            "type": log.removal_type,
            "status": log.status,
            "message": log.message,
        })

    activities.sort(
        key=lambda x: x["date_time"],
        reverse=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Activity Logs"

    headers = [
        "Date / Time",
        "User",
        "Action",
        "Code",
        "Reference",
        "Type",
        "Status",
        "Message",
    ]

    worksheet.append(headers)

    for activity in activities:
        worksheet.append([
            activity["date_time"].replace(tzinfo=None),
            activity["user"],
            activity["action"],
            activity["code"],
            activity["reference"],
            activity["type"],
            activity["status"],
            activity["message"],
        ])

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="activity_logs.xlsx"'
    )

    workbook.save(response)

    return response
